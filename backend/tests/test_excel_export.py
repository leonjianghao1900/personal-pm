"""Tests for Excel export endpoint."""
from openpyxl import load_workbook

from app.models import Cost, Task


def test_export_returns_xlsx(client):
    pid = client.post("/api/projects", json={"name": "测试项目"}).json()["id"]
    r = client.get(f"/api/projects/{pid}/export")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert r.headers["content-disposition"].startswith("attachment")
    body = r.content
    assert body[:2] == b"PK"  # xlsx is a zip file


def test_export_has_four_sheets(client):
    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    r = client.get(f"/api/projects/{pid}/export")
    wb = load_workbook(__import__("io").BytesIO(r.content))
    assert wb.sheetnames == ["项目概览", "任务清单", "费用记录", "甘特图"]


def test_export_includes_tasks_and_costs(client, db_session):
    from datetime import date

    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    t = Task(project_id=pid, name="任务1", status="progress", start_date=date(2026, 8, 20), end_date=date(2026, 8, 25), progress=50)
    db_session.add(t)
    db_session.flush()
    c = Cost(project_id=pid, task_id=t.id, name="材料", category="物料", amount=123.45, spent_at=date(2026, 8, 21))
    db_session.add(c)
    db_session.commit()

    r = client.get(f"/api/projects/{pid}/export")
    wb = load_workbook(__import__("io").BytesIO(r.content))

    tasks_sheet = wb["任务清单"]
    task_rows = [row for row in tasks_sheet.iter_rows(min_row=2, values_only=True) if row[1] == "任务1"]
    assert len(task_rows) == 1
    assert task_rows[0][3] == "进行中"
    assert task_rows[0][7] == 6  # duration
    assert task_rows[0][8] == "50%"

    costs_sheet = wb["费用记录"]
    cost_rows = [row for row in costs_sheet.iter_rows(min_row=2, values_only=True) if row[1] == "材料"]
    assert len(cost_rows) == 1
    assert cost_rows[0][2] == "物料"
    assert cost_rows[0][3] == 123.45

    gantt_sheet = wb["甘特图"]
    gantt_rows = [row for row in gantt_sheet.iter_rows(min_row=2, values_only=True) if row[1] == "任务1"]
    assert len(gantt_rows) == 1
    assert gantt_rows[0][2] == "2026-08-20"


def test_export_nonexistent_returns_404(client):
    r = client.get("/api/projects/9999/export")
    assert r.status_code == 404


def test_export_handles_project_without_dates(client):
    pid = client.post("/api/projects", json={"name": "NoDates"}).json()["id"]
    r = client.get(f"/api/projects/{pid}/export")
    assert r.status_code == 200
    wb = load_workbook(__import__("io").BytesIO(r.content))
    assert wb["甘特图"].max_row == 1  # only header
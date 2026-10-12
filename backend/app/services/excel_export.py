"""Excel export service: generate .xlsx with project, tasks, costs, and Gantt PNG."""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont
from sqlalchemy.orm import Session

from app.models import Cost, Project, Task


STATUS_LABELS = {
    "todo": "未开始",
    "progress": "进行中",
    "delayed": "延期",
    "done": "已完成",
    "cancelled": "已取消",
}
PRIORITY_LABELS = {"low": "低", "normal": "中", "high": "高"}
STATUS_COLORS = {
    "todo": "#8e8e93",
    "progress": "#0071e3",
    "delayed": "#ff9500",
    "done": "#34c759",
    "cancelled": "#af52de",
}

THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
HEADER_FILL = PatternFill("solid", fgColor="23282D")
HEADER_FONT = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Microsoft YaHei", size=10)


def _hex_to_argb(hex_color: str) -> tuple[int, int, int]:
    h = hex_color.lstrip("#")
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def _render_gantt_png(project: Project, tasks: list[Task]) -> bytes:
    """Render a PNG Gantt chart using Pillow (no matplotlib)."""
    dated = [t for t in tasks if t.start_date]
    if not dated:
        return _render_placeholder_png("暂无带日期的任务")

    min_d = min(t.start_date for t in dated)
    max_d = max(t.end_date or t.start_date for t in dated)
    if max_d < min_d:
        max_d = min_d
    min_d = min_d - timedelta(days=1)
    max_d = max_d + timedelta(days=1)
    total_days = (max_d - min_d).days + 1

    name_col_w = 200
    day_w = max(20, min(60, 1100 // total_days))
    width = name_col_w + day_w * total_days + 40
    row_h = 32
    header_h = 70
    height = header_h + row_h * len(dated) + 40

    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)

    try:
        title_font = ImageFont.truetype("msyh.ttc", 16)
        small_font = ImageFont.truetype("msyh.ttc", 11)
        day_font = ImageFont.truetype("msyh.ttc", 9)
    except OSError:
        title_font = ImageFont.load_default()
        small_font = ImageFont.load_default()
        day_font = ImageFont.load_default()

    draw.text((20, 12), f"项目：{project.name}", fill="#1d1d1f", font=title_font)
    draw.text((20, 38), f"时间范围：{min_d.isoformat()} → {max_d.isoformat()}",
              fill="#6e6e73", font=small_font)

    # Day grid header
    grid_y = header_h - 10
    draw.rectangle([20, grid_y, width - 20, grid_y + 22], fill="#e8e8eb", outline="#d0d0d0")
    draw.text((24, grid_y + 4), "任务", fill="#1d1d1f", font=small_font)
    for d in range(total_days):
        x = 20 + name_col_w + d * day_w
        date_obj = min_d + timedelta(days=d)
        label = date_obj.strftime("%m-%d")
        text_color = "#0071e3" if date_obj == date.today() else "#1d1d1f"
        draw.text((x + 4, grid_y + 4), label, fill=text_color, font=day_font)
        if date_obj == date.today():
            draw.line([(x + day_w // 2, grid_y + 18), (x + day_w // 2, height - 20)], fill="#0071e3", width=1)

    # Rows
    for i, t in enumerate(dated):
        y = header_h + 12 + i * row_h
        if i % 2 == 0:
            draw.rectangle([20, y, width - 20, y + row_h - 4], fill="#f8f8fa")
        draw.text((24, y + 8), t.name[:24], fill="#1d1d1f", font=small_font)
        start_d = t.start_date
        end_d = t.end_date or t.start_date
        start_col = (start_d - min_d).days
        end_col = (end_d - min_d).days
        bar_x1 = 20 + name_col_w + start_col * day_w + 2
        bar_x2 = 20 + name_col_w + (end_col + 1) * day_w - 2
        rgb = _hex_to_argb(STATUS_COLORS.get(t.status, "#8e8e93"))
        draw.rectangle([bar_x1, y + 6, bar_x2, y + row_h - 12], fill=rgb, outline=None)

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _render_placeholder_png(text: str) -> bytes:
    img = Image.new("RGB", (800, 200), "white")
    draw = ImageDraw.Draw(img)
    try:
        f = ImageFont.truetype("msyh.ttc", 18)
    except OSError:
        f = ImageFont.load_default()
    draw.text((20, 80), text, fill="#8e8e93", font=f)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _summary(db: Session, project: Project) -> dict:
    tasks = list(project.tasks)
    costs = list(project.costs)
    total = sum(c.amount for c in costs)
    today = date.today()
    month_key = today.strftime("%Y-%m")
    this_month = sum(c.amount for c in costs if c.spent_at.strftime("%Y-%m") == month_key)
    done = sum(1 for t in tasks if t.status == "done")
    delayed = sum(1 for t in tasks if t.status == "delayed")
    avg_progress = sum(t.progress for t in tasks) / len(tasks) if tasks else 0
    return {
        "task_count": len(tasks),
        "done": done,
        "delayed": delayed,
        "avg_progress": avg_progress,
        "total_cost": total,
        "this_month_cost": this_month,
    }


def _write_overview(wb: Workbook, project: Project, summary: dict) -> None:
    ws = wb.active
    ws.title = "项目概览"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 70

    rows = [
        ("名称", project.name),
        ("颜色", project.color),
        ("状态", project.status),
        ("开始日期", project.start_date.isoformat() if project.start_date else ""),
        ("结束日期", project.end_date.isoformat() if project.end_date else ""),
        ("描述", project.description or ""),
        ("", ""),
        ("任务总数", summary["task_count"]),
        ("已完成", summary["done"]),
        ("延期", summary["delayed"]),
        ("平均进度", f"{summary['avg_progress']:.0f}%"),
        ("本月支出", f"¥{summary['this_month_cost']:.2f}"),
        ("总支出", f"¥{summary['total_cost']:.2f}"),
        ("", ""),
        ("导出时间", datetime.utcnow().isoformat(timespec="seconds")),
    ]
    for r, (k, v) in enumerate(rows, start=1):
        cell_k = ws.cell(r, 1, k)
        cell_v = ws.cell(r, 2, v)
        if k:
            cell_k.font = Font(name="Microsoft YaHei", size=10, bold=True)
            cell_k.fill = PatternFill("solid", fgColor="F2F2F5")
        cell_v.font = BODY_FONT
        cell_v.alignment = Alignment(vertical="top", wrap_text=True)
        if k == "描述":
            ws.row_dimensions[r].height = 60


def _write_tasks(wb: Workbook, project: Project, tasks_by_id: dict[int, Task]) -> None:
    ws = wb.create_sheet("任务清单")
    headers = ["ID", "名称", "父任务", "状态", "优先级", "开始", "结束", "持续(天)", "进度", "里程碑", "描述"]
    widths = [6, 30, 8, 10, 8, 12, 12, 10, 8, 8, 40]
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(1, c, h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"

    sorted_tasks = sorted(project.tasks, key=lambda t: (t.parent_id or 0, t.sort_order, t.id))
    for r, t in enumerate(sorted_tasks, start=2):
        parent_name = tasks_by_id.get(t.parent_id).name if t.parent_id and t.parent_id in tasks_by_id else ""
        days = ""
        if t.start_date and t.end_date:
            days = (t.end_date - t.start_date).days + 1
        values = [
            t.id,
            t.name,
            parent_name,
            STATUS_LABELS.get(t.status, t.status),
            PRIORITY_LABELS.get(t.priority, t.priority),
            t.start_date.isoformat() if t.start_date else "",
            t.end_date.isoformat() if t.end_date else "",
            days,
            f"{t.progress}%",
            "是" if t.is_milestone else "",
            t.description or "",
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(r, c, v)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(c == 11))
            if c == 4 and t.status in STATUS_COLORS:
                rgb = _hex_to_argb(STATUS_COLORS[t.status])
                cell.font = Font(name="Microsoft YaHei", size=10, color="{:02X}{:02X}{:02X}".format(*rgb))
        if t.description:
            ws.row_dimensions[r].height = 30


def _write_costs(wb: Workbook, project: Project) -> None:
    ws = wb.create_sheet("费用记录")
    headers = ["ID", "名称", "类别", "金额", "日期", "关联任务", "描述"]
    widths = [6, 28, 10, 12, 12, 10, 36]
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(1, c, h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"

    sorted_costs = sorted(project.costs, key=lambda c: (c.spent_at, c.id))
    for r, c in enumerate(sorted_costs, start=2):
        values = [
            c.id,
            c.name,
            c.category,
            c.amount,
            c.spent_at.isoformat(),
            f"#{c.task_id}" if c.task_id else "",
            c.description or "",
        ]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(r, col, v)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(col == 7))
            if col == 4:
                cell.number_format = "¥#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")


def _write_gantt(wb: Workbook, project: Project, tasks: list[Task], png_bytes: bytes) -> None:
    ws = wb.create_sheet("甘特图")
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 8

    headers = ["ID", "任务名", "开始", "结束", "状态", "进度"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(1, c, h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER

    dated = sorted([t for t in tasks if t.start_date], key=lambda t: t.start_date)
    for r, t in enumerate(dated, start=2):
        values = [
            t.id,
            t.name,
            t.start_date.isoformat() if t.start_date else "",
            t.end_date.isoformat() if t.end_date else "",
            STATUS_LABELS.get(t.status, t.status),
            f"{t.progress}%",
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(r, c, v)
            cell.font = BODY_FONT
            cell.border = BORDER
            if c == 5 and t.status in STATUS_COLORS:
                rgb = _hex_to_argb(STATUS_COLORS[t.status])
                cell.font = Font(name="Microsoft YaHei", size=10, color="{:02X}{:02X}{:02X}".format(*rgb))

    from openpyxl.drawing.image import Image as XLImg
    img_row = len(dated) + 4
    img = XLImg(io.BytesIO(png_bytes))
    img.width = 1100
    img.height = max(120, min(700, 32 * len(dated) + 80))
    img.anchor = f"A{img_row}"
    ws.add_image(img)


def export_project_workbook(db: Session, project_id: int) -> bytes:
    project = db.get(Project, project_id)
    if not project:
        raise ValueError(f"project {project_id} not found")
    db.refresh(project)
    tasks_by_id = {t.id: t for t in project.tasks}
    summary = _summary(db, project)

    wb = Workbook()
    _write_overview(wb, project, summary)
    _write_tasks(wb, project, tasks_by_id)
    _write_costs(wb, project)
    png_bytes = _render_gantt_png(project, list(project.tasks))
    _write_gantt(wb, project, list(project.tasks), png_bytes)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()